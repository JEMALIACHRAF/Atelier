import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
E = sys.argv[1]
NAVY = RGBColor(0, 0x3a, 0x70); ORG = RGBColor(0xc7, 0x6b, 0x2b); INK = RGBColor(0x1a, 0x23, 0x32)
MUTE = RGBColor(0x8a, 0x94, 0xa6); RULE = RGBColor(0xe2, 0xe5, 0xea); PAPER = RGBColor(0xfa, 0xfb, 0xfc)
WHITE = RGBColor(255, 255, 255); CRIT = RGBColor(0xb9, 0x38, 0x38); OK = RGBColor(0x2e, 0x7d, 0x4f)
SER = "Source Serif Pro"; SAN = "Source Sans Pro"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def tb(s, x, y, w, h, text, size=14, bold=False, color=INK, font=SAN, align=None, anchor=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f = t.text_frame; f.word_wrap = True
    if anchor: f.vertical_anchor = anchor
    paras = text if isinstance(text, list) else [text]
    for i, p in enumerate(paras):
        para = f.paragraphs[0] if i == 0 else f.add_paragraph()
        r = para.add_run(); r.text = p; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
        para.space_after = Pt(6)
        if align: para.alignment = align
    return t


def bullets(s, x, y, w, h, items, size=14):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f = t.text_frame; f.word_wrap = True
    for i, it in enumerate(items):
        para = f.paragraphs[0] if i == 0 else f.add_paragraph()
        head = None
        if isinstance(it, tuple): head, it = it
        r0 = para.add_run(); r0.text = "• "; r0.font.size = Pt(size); r0.font.color.rgb = ORG; r0.font.name = SAN; r0.font.bold = True
        if head:
            r = para.add_run(); r.text = head + " "; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY; r.font.name = SAN
        r = para.add_run(); r.text = it; r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = SAN
        para.space_after = Pt(8)


def rect(s, x, y, w, h, fill, line=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); r.fill.solid(); r.fill.fore_color.rgb = fill
    if line: r.line.color.rgb = line
    else: r.line.fill.background()
    r.shadow.inherit = False
    return r


n = [0]


def slide(title, source):
    s = prs.slides.add_slide(blank); n[0] += 1
    rect(s, 0, 0, 13.333, 0.12, NAVY)
    tb(s, 0.6, 0.35, 12.1, 1.2, title, 24, True, NAVY, SER, anchor=MSO_ANCHOR.TOP)
    rect(s, 0.6, 1.55, 1.2, 0.05, ORG)
    ln = s.shapes.add_connector(1, Inches(0.6), Inches(6.85), Inches(12.73), Inches(6.85)); ln.line.color.rgb = RULE
    tb(s, 0.6, 6.9, 11.3, 0.5, source, 9, False, MUTE)
    tb(s, 12.0, 6.9, 0.8, 0.4, str(n[0]), 10, False, MUTE, align=PP_ALIGN.RIGHT)
    return s


def callout(s, x, y, w, h, head, body, color=ORG, size=12):
    rect(s, x, y, w, h, PAPER, RULE); rect(s, x, y, 0.08, h, color)
    tb(s, x + 0.2, y + 0.1, w - 0.3, 0.4, head, 12, True, color)
    tb(s, x + 0.2, y + 0.5, w - 0.3, h - 0.55, body, size, False, INK)


def table(s, x, y, w, rows, colw, size=11, rowh=0.4):
    sh = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))); t = sh.table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = v; r.font.size = Pt(size); r.font.name = SAN
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = NAVY; r.font.color.rgb = WHITE; r.font.bold = True
            else:
                c.fill.fore_color.rgb = WHITE if i % 2 else PAPER; r.font.color.rgb = INK
                if j == 0: r.font.bold = True; r.font.color.rgb = NAVY
    return t


LIGHT = RGBColor(0xc9, 0xd6, 0xe6)
# 1 cover + these
s = prs.slides.add_slide(blank); n[0] += 1
rect(s, 0, 0, 13.333, 7.5, NAVY); rect(s, 0.7, 2.05, 1.2, 0.06, ORG)
tb(s, 0.7, 3.2, 11.9, 0.45, "Un seul signal d'avance, la souverainete, existe au niveau groupe et reste non conclusif [D22, D24].", 13, True, RGBColor(0xf0, 0xb0, 0x80))
tb(s, 0.7, 0.5, 11, 0.4, "ATELIER  |  Benchmark concurrentiel  |  Pre-brief express  |  24 octobre 2026", 12, False, LIGHT)
tb(s, 0.7, 0.95, 11.9, 1.1, "Sur donnees publiques, Invent n'a aucune avance IA etablie en Europe ; son vrai defi est d'etre visible", 27, True, WHITE, SER)
tb(s, 0.7, 2.3, 11.9, 0.9, "Positionnement de Capgemini Invent face a McKinsey QuantumBlack, BCG X, Accenture et Deloitte (2024-2026, sources publiques). Pre-brief express : les lacunes de donnees sont affichees, pas comblees.", 13, False, LIGHT)
msgs = [("1. Pas d'avance etablie", "Groupe Capgemini a parite sur l'agentique (3 vs 3 pour Accenture, Deloitte et McKinsey ; BCG X 2) ; 3 dimensions sur 4 non conclusives [D01-D06, D26, D61-D69]."),
        ("2. Ecart d'attribution et d'echelle", "Aucun actif IA public n'est rattache a Invent [D66]. Retard robuste d'echelle vs Accenture : 77 000 professionnels data & IA [D30, S15]."),
        ("3. Angle 2027 = hypothese a valider", "IA agentique souveraine europeenne : 2 partenaires souverains + 1 offre combinee au niveau groupe [D22, non verifie au primaire] vs 1 chez Accenture [D23] ; McKinsey liste Mistral et Aleph Alpha [D24, D63] ; rien n'est attribue a Invent [D66].")]
for i, (h, b) in enumerate(msgs):
    x = 0.7 + i * 4.1; rect(s, x, 3.75, 3.85, 2.75, RGBColor(0, 0x2a, 0x52)); rect(s, x, 3.75, 3.85, 0.07, ORG)
    tb(s, x + 0.2, 3.95, 3.5, 0.5, h, 16, True, WHITE, SER); tb(s, x + 0.2, 4.55, 3.5, 1.9, b, 11.5, False, RGBColor(0xe2, 0xe8, 0xf0))
tb(s, 0.7, 6.6, 12, 0.5, "Sources : 02_data/sources.json (67 sources, 8 verifiees au primaire) ; benchmark v2. Chiffres au niveau groupe et mondial sauf mention.", 10, False, RGBColor(0xa9, 0xb8, 0xcc))

# 2
s = slide("Au niveau groupe, Capgemini est a parite sur l'agentique ; 3 dimensions sur 4 restent non conclusives",
          "Source : benchmark v2, scoring_matrix.json ; D01-D06, D30, D59-D69 ; S1, S6, S15, S21, S23, S30, S31. Echelle 1-5, 2024-2026, perimetre groupe/mondial. Gris = ND/NC, pas faiblesse.")
s.shapes.add_picture(f"{E}/03_charts/scoring_heatmap.png", Inches(0.6), Inches(1.8), width=Inches(8.0))
bullets(s, 8.9, 1.8, 3.9, 3.2, [("Bilan :", "0 avance ; 0 retard ferme ; 1 retard robuste (D4 vs Accenture) ; 1 retard provisoire (D1 Invent strict) ; 1 parite (D1 groupe) ; 3 non conclusives."),
                                 ("Cible du brief non atteinte :", "2-3 avances / 2-3 retards non forces. Lire ND comme un score violerait la regle du workplan.")], 12.5)
callout(s, 8.9, 5.0, 3.9, 1.7, "Contre-indice", "Accenture probablement sous-note en D1 (references Ecolab/UOB non europeennes) [D07, S15].", ORG, 11.5)

# 3
s = slide("L'agentique est devenue un prerequis : 4 acteurs sur 5 (Capgemini compte au niveau groupe) ont une offre nommee, aucun n'atteint le niveau 4 ; Invent strict n'en a aucune attribuee",
          "Source : D01-D07, D10, D66 ; S1, S6, S21, S23, S30, S31. Niveaux 1-5 (rubrique du workplan), 2024-2026, perimetre groupe/mondial. DP02 (references clients) : data not found.")
table(s, 0.6, 1.85, 7.6, [["Acteur", "Offre agentique publique", "Score D1"],
                          ["Capgemini groupe", "Agentique industrielle NVIDIA, CX agentique Google [D01, D02]", "3"],
                          ["Accenture", "AI Refinery for Industry, 12 solutions d'agents [D03]", "3"],
                          ["Deloitte", "Zora AI, Global Agentic Network [D06]", "3"],
                          ["McKinsey QuantumBlack", "Suite \"agents at scale\" [D04]", "3"],
                          ["BCG X", "Offre moins documentee [D05]", "2"],
                          ["Capgemini Invent (strict)", "Aucune attribution publique [D66]", "1* (plancher)"]], [2.2, 4.0, 1.4], 11.5, 0.55)
bullets(s, 8.6, 1.85, 4.2, 3.0, [("So what :", "\"Nous avons des offres agentiques\" ne distingue Invent d'aucun concurrent."),
                                  ("Niveau 4 :", "aucun acteur, faute de references clients publiques comptees [D07 : data not found]."),
                                  ("*", "Le 1 d'Invent mesure l'absence d'attribution, pas une faiblesse mesuree.")], 13)
callout(s, 8.6, 5.0, 4.2, 1.7, "Divergence de donnees", "Agents McKinsey : 20 000 vs 25 000 selon la source [D10].", ORG, 12)

# 4
s = slide("Sur les hyperscalers, rien ne distingue Capgemini : parite indicative avec Deloitte, niveaux de partenariat des concurrents non publies",
          "Source : D18, D59-D66 ; S10, S53, S55, S56 (verifiees par fetch : S53, S55, S56). Tiers partenaires, 2024-2026, mondial. Annuaires AWS/Microsoft/Google non interroges (D65).")
table(s, 0.6, 1.85, 7.8, [["Acteur", "Statut hyperscalers", "Statut"],
                          ["Capgemini groupe", "Top tier AWS [S53] et Google [S10] ; Microsoft ND [D59, D60]", "Indicatif 3"],
                          ["Deloitte", "Egale Capgemini chez AWS [D62, S56]", "Indicatif 3"],
                          ["Accenture", "Awards (20x Microsoft SI of the Year, 4x Google) [D61, S55] ; tiers ND", "NC"],
                          ["McKinsey", "Alliances avec les 3 hyperscalers, sans termes [D63]", "NC"],
                          ["BCG X", "Tiers ND [D64]", "NC"],
                          ["Invent (strict)", "Aucun statut attribue [D66]", "NC"]], [1.9, 4.6, 1.3], 11.5, 0.58)
bullets(s, 8.8, 1.85, 4.0, 3.2, [("Non conclusif :", "4 acteurs sur 6 en data not found, donc pas de lecture avance/retard."),
                                  ("A noter :", "Accenture a le portefeuille de modeles de fondation le plus large (Anthropic, OpenAI, NVIDIA, Mistral) [D18]."),
                                  ("H2.1 :", "non testable ; signal faible contre une avance sur McKinsey.")], 13)

# 5
s = slide("McKinsey est le seul des 5 confirme partenaire VivaTech 2025, Capgemini en est absent ; conferences et reprises presse restent non mesurees",
          "Source : D26, D49-D58 ; S40 (fetch), S66. Volumes = bornes basses non comparables [D54]. 2024-2026, mondial. D3 = publications + conferences fusionnees (scope express).")
table(s, 0.6, 1.85, 6.6, [["Acteur", "Publications phares (borne basse)", "Conferences"],
                          ["Capgemini (CRI)", "4-5 [D49-D53]", "Absent des tiers VivaTech 2025 [D56]"],
                          ["Deloitte", "4", "ND / indirect"],
                          ["Accenture", "3", "ND / indirect"],
                          ["McKinsey", "1", "Gold Partner VivaTech 2025 [D55, S40]"],
                          ["BCG X", "1", "ND / indirect"]], [1.8, 2.3, 2.5], 11.5, 0.6)
bullets(s, 7.6, 1.85, 5.2, 3.4, [("Non conclusif :", "reprises presse ND pour les 5 acteurs [D26] ; conferences ND ou indirectes pour 4 sur 5 [D55-D58]."),
                                  ("Attribution :", "l'editeur est le Capgemini Research Institute, pas Invent [S66]."),
                                  ("H3.1 :", "non testable. Les volumes ne sont pas comparables (think tank, practice, entite) [D54].")], 13)
callout(s, 0.6, 5.55, 6.6, 1.15, "Piste a reintegrer", "PwC, EY et KPMG etaient sponsors de VivaTech 2025 [D55, S40] : argument pour la reintegration de PwC en mode standard.", ORG, 11.5)

# 6
s = slide("Face aux 77 000 professionnels data & IA d'Accenture (definition large, groupe, monde) [D30], Invent ne peut pas jouer l'echelle",
          "Source : D30 (S15, fetch), D33 (S60), D34, D35 (S11), D37-D38, D67-D69. Effectifs 2021-2026, perimetres heterogenes (practice vs entite). D68 et D69 : data not found.")
table(s, 0.6, 1.85, 7.4, [["Acteur", "Effectif IA/data", "Niveau"],
                          ["Accenture (Data & AI)", "77 000 pros AI & data [D30, S15]", "5 (definition large)"],
                          ["BCG X", "3 000+ [D34]", "2"],
                          ["QuantumBlack", "~1 700 [D33] (~7 000 selon perimetre, S60)", "2"],
                          ["Invent (strict)", "\"10 000+\" toutes fonctions, texte de 2021 [D35, S11]", "<= 3 (ND)"],
                          ["Capgemini groupe", "Effectif data & IA non publie [D68]", "ND"],
                          ["Deloitte", "ND", "ND"]], [2.2, 3.7, 1.5], 11.5, 0.55)
bullets(s, 8.4, 1.85, 4.4, 3.6, [("Robuste :", "Invent <= 3 vs Accenture 5, quelle que soit la valeur reelle d'Invent."),
                                  ("Indetermine :", "vs QuantumBlack et BCG X (1-3 vs 2)."),
                                  ("A ne pas confondre :", "421 000 collaborateurs [D67] et 310 000 formes a la GenAI [D38] ne sont pas des effectifs IA."),
                                  ("So what :", "angle de specialisation ou d'ancrage, pas de volume.")], 12.5)

# 7
s = slide("Aucun actif IA public ne porte le nom Invent : c'est le probleme de positionnement numero 1",
          "Source : D66 (2 recherches ciblees), D47 ; S4 (fetch), S66, S67 ; D33, D34, D04, D05, D55. Constat sur sources publiques 2024-2026 ; a reverifier au primaire.")
bullets(s, 0.6, 1.85, 6.9, 4.8, [("Agentique NVIDIA/Google, Mistral, Bleu, etude \"Rise of Agentic AI\" :", "tous attribues au groupe ou au CRI, jamais a Invent [D66, S66]."),
                                  ("Leadership agentique :", "affiche sous Insights & Data [S67]."),
                                  ("Capital Markets Day 2026 :", "Invent n'apparait que comme une marque [D47, S4]."),
                                  ("Aucun tier hyperscaler [D59] ni effectif IA [D69]", "n'est attribue a Invent.")], 14)
callout(s, 7.9, 1.85, 4.9, 2.0, "A l'inverse : QuantumBlack et BCG X", "Marques IA autonomes, avec effectif [D33, D34], produits [D04, D05] et visibilite propres (Gold Partner VivaTech 2025 [D55]).", NAVY, 12.5)
callout(s, 7.9, 4.1, 4.9, 2.5, "So what", "Pour un acheteur ou analyste, Invent n'a pas d'identite IA. Le rattachement des actifs groupe a Invent est le prealable a tout angle 2027. Ecart d'attribution, pas ecart de capacite demontre.", ORG, 12.5)

# 8
s = slide("Angle 2027 a valider : l'IA agentique souveraine europeenne, a condition de l'attribuer a Invent avant qu'Accenture ne le referme",
          "Source : D22-D24, D42, D45-D47, D63, D66 ; S4, S7-S9, S18, S19, S35, S37. D45 = etude commanditee par Accenture ; D22 (S7-S9) et D42 (S37) non verifies au primaire. DP17 : data not found. Hypothese, pas recommandation ferme.")
tb(s, 0.6, 1.75, 12.1, 0.5, "Hypothese : \"Capgemini Invent, conseil de reference pour l'IA agentique souveraine et conforme en Europe\"", 14, True, INK, SER)
callout(s, 0.6, 2.35, 4.0, 3.0, "Pour", "Groupe : 2 partenaires souverains distincts (Mistral, Bleu) + 1 offre combinee SAP-Bleu-Mistral [D22, NON VERIFIE au primaire] vs 1 partenariat Accenture [D23]. Priorite au CMD 2026 [D47]. Demande : 62 % [D45], 61 % [D46], 48 % citent la conformite [D45].", OK, 11.5)
callout(s, 4.75, 2.35, 4.0, 3.1, "Contre", "Accord Accenture-Mistral, fev. 2026 [D23, S18]. McKinsey liste Mistral et Aleph Alpha (listing, pas accord) [D24, D63]. 62 % = etude Accenture ; 65 % ne peuvent pas se passer de non-europeens ; pas d'etude independante sur l'AI Act [D45, DP17]. McKinsey ~1/3 outcome-based [D42, NON VERIFIE au primaire]. Non attribue a Invent [D66].", CRIT, 10)
callout(s, 8.9, 2.35, 3.9, 3.0, "Go / no-go (4 tests)", "1. >= 2 actifs attribues a Invent. 2. >= 3 references clients europeennes. 3. Concurrents <= 1 partenariat souverain (DP07). 4. >= 1 etude independante sur la demande. No-go si test 1 echoue a 6 mois.", NAVY, 11)
table(s, 0.6, 5.55, 12.2, [["Horizon", "Owner", "KPI"],
                           ["Validation T1-T2 2027 ; lancement S2 2027 au plus tot si go", "CEO Capgemini Invent, avec Data & AI groupe et marketing d'offres", "Actifs attribues a Invent : 0 [D66] -> >= 2 a 6 mois ; references clients europeennes >= 3 a 12 mois"]], [3.7, 3.8, 4.7], 10.5, 0.55)

# 9
s = slide("4 decisions pour Invent : rendre ses actifs IA visibles d'abord, valider l'angle souverain ensuite",
          "Source : 05_synthesis.md v2 (Recommandations) ; D66, D30, D42 (non verifie au primaire). Detail des lacunes de donnees : annexe A1 (slide 10).")
table(s, 0.6, 1.85, 12.2, [["#", "Decision", "Horizon", "Owner", "KPI"],
   ["1", "Lancer un chantier d'attribution : rattacher publiquement a Invent des actifs IA du groupe (offre, partenariat ou etude)", "6 mois", "Direction de Capgemini Invent + marketing d'offres", "Actifs IA publiquement attribues a Invent : baseline 0 [D66], cible >= 2"],
   ["2", "Soumettre l'angle \"IA agentique souveraine europeenne\" au test go/no-go avant de l'adopter", "T1-T2 2027", "CEO d'Invent", "4 tests valides ; no-go si l'attribution echoue a 6 mois ou si >= 2 concurrents ont >= 2 partenariats souverains equivalents"],
   ["3", "Ne pas construire le positionnement 2027 sur l'echelle ni sur le seul outcome-based [D30 ; D42, non verifie au primaire]", "Immediat (cadrage du message 2027)", "Direction d'Invent", "Aucun claim d'echelle dans les supports Invent 2027"],
   ["4", "Relancer l'etude en mode standard avant toute communication de positionnement", "4 semaines", "Equipe engagement Atelier (sponsor : direction d'Invent)", "D2, D3, D4 passees en conclusive ; >= 80 % des S# cites verifies au primaire (lacunes : annexe A1)"]],
   [0.4, 4.6, 1.8, 2.3, 3.1], 10.5, 0.8)
callout(s, 0.6, 6.1, 12.2, 0.7, "Ordre de priorite", "L'attribution passe avant l'angle : sans rattachement des actifs groupe a Invent, aucun angle 2027 ne profite a Invent.", ORG, 11.5)

# 10
s = slide("Annexe A1 : methodologie et 5 lacunes de donnees - perimetre express reduit, 59 sources sur 67 non verifiees au primaire",
          "Source : decisions 2026-10-23 (scope-express), 2026-10-24 (prebrief-honnete) ; 02_data/sources.json ; D26, D37, D55-D58, D65-D69, DP02, DP04, DP07, DP09. Bibliographie complete (67 sources, URL) dans report.docx.")
bullets(s, 0.6, 1.8, 5.6, 5.0, [("Perimetre reduit :", "4 concurrents ; PwC AI CoE retire (plafond express, perimetre flou) ; EY, KPMG, IBM hors perimetre."),
    ("D3 fusionnee :", "publications + conferences en une dimension (plafond de 4), sous-scores lacunaires."),
    ("Notation :", "echelle 1-5 ; M = mediane des 4 concurrents ; non conclusive si plus de 2 acteurs en ND ; ND != faible. Override D4 valide par le Chief."),
    ("Sources :", "8 sur 67 verifiees par fetch (S2, S4, S15, S19, S40, S53, S55, S56) ; 59 non verifiees au primaire, dont S7-S9 (D22) et S37 (D42) qui portent l'angle 2027."),
    ("Geographie :", "groupe et mondial ; vue Europe quasi absente (D37). 2024-2026, sources publiques."),
    ("Hypotheses :", "8 sur 8 traitees, aucune confirmee ni refutee.")], 11)
table(s, 6.4, 1.85, 6.4, [["#", "Lacune de donnees (mode standard)", "Source a interroger"],
   ["1", "Rattachement a Invent [D66, D69]", "capgemini.com/invent, communiques, organigramme Data & AI"],
   ["2", "Tiers hyperscalers [D65]", "AWS Partner Finder, Microsoft Partner Finder, Google Cloud Partner Advantage"],
   ["3", "Effectifs data & IA, repartition Europe [D68, D69, D37]", "Document d'enregistrement universel, transcript CMD, Deloitte AI & Data"],
   ["4", "References clients agentiques ; partenariats souverains concurrents [DP02, DP04, DP07, D63]", "Comptage total / Europe / industrie ; qualifier les listings McKinsey"],
   ["5", "D3 et PwC [D26, D55-D58, DP09]", "Programmes speakers, reprises presse ; reintegrer PwC ; verifier S7-S9, S37"]], [0.4, 3.0, 3.0], 10, 0.75)
prs.save(f"{E}/06_deliverables/deck.pptx")
print("saved", n[0])
