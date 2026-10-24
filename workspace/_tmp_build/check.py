import os, sys
from pptx import Presentation
from docx import Document
E = sys.argv[1]
p = Presentation(E + '/06_deliverables/deck.pptx'); print('slides', len(p.slides))
d = Document(E + '/06_deliverables/report.docx')
words = sum(len(x.text.split()) for x in d.paragraphs) + sum(len(c.text.split()) for t in d.tables for r in t.rows for c in r.cells)
print('paras', len(d.paragraphs), 'tables', len(d.tables), 'images', len(d.inline_shapes), 'words', words)
for f in os.listdir(E + '/06_deliverables'): print(f, os.path.getsize(E + '/06_deliverables/' + f))
for n in ('soffice.exe',):
    print(os.path.exists('C:/Program Files/LibreOffice/program/' + n))
