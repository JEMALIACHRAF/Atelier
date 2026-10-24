import sys, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
E = sys.argv[1]
NAVY = RGBColor(0, 0x3a, 0x70); ORG = RGBColor(0xc7, 0x6b, 0x2b)
d = Document()
sec = d.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.0)
st = d.styles["Normal"]; st.font.name = "Source Sans Pro"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Source Sans Pro")
st.paragraph_format.space_after = Pt(6)
for name, sz in (("Heading 1", 16), ("Heading 2", 12.5)):
    h = d.styles[name]; h.font.name = "Source Serif Pro"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = NAVY
    h.element.rPr.rFonts.set(qn("w:ascii"), "Source Serif Pro"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Source Serif Pro")
    h.paragraph_format.space_before = Pt(14 if sz > 13 else 10); h.paragraph_format.space_after = Pt(6)


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


def para(text, bold=False, italic=False, size=None, color=None, align=None, style=None):
    p = d.add_paragraph(style=style)
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    return p


def rich(parts, style=None):
    p = d.add_paragraph(style=style)
    for t, b in parts:
        r = p.add_run(t); r.bold = b
    return p


def bl(items):
    for it in items:
        if isinstance(it, tuple): rich([(it[0] + " ", True), (it[1], False)], "List Bullet")
        else: d.add_paragraph(it, style="List Bullet")


def table(rows, widths, size=9):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(size)
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "003A70")
            elif j == 0: r.bold = True
    d.add_paragraph()
    return t


def callout(head, body):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.cell(0, 0); shade(c, "FAFBFC")
    p = c.paragraphs[0]; r = p.add_run(head + " "); r.bold = True; r.font.color.rgb = ORG
    c.add_paragraph(body); d.add_paragraph()


# cover
for _ in range(6): d.add_paragraph()
para("ATELIER  |  Benchmark concurrentiel  |  Pre-brief express", size=11, color=ORG, bold=True)
para("Positionnement de Capgemini Invent sur l'IA en Europe face a QuantumBlack, BCG X, Accenture et Deloitte", bold=True, size=26, color=NAVY)
para("Sur donnees publiques, Invent n'a aucune avance IA etablie en Europe ; son vrai defi est d'etre visible.", italic=True, size=13)
d.add_paragraph()
para("Rapport court, horizon d'observation 2024-2026, sources publiques uniquement.\nMode express : pre-brief, pas evaluation finale.\n24 octobre 2026", size=11)
d.add_page_break()

# 1 exec summary
d.add_heading("1. Synthese executive", 1)
callout("These centrale.", "Sur donnees publiques, aucune avance de Capgemini Invent n'est etablie. Le seul ecart robuste est un retard d'echelle des equipes IA face a Accenture. Les actifs IA visibles (agentique NVIDIA/Google, Mistral, Bleu, CRI) sont tous attribues au groupe Capgemini et jamais a Invent : cette invisibilite est en soi un probleme de positionnement [D66, S66, S67].")
para("Un seul signal d'avance apparait, la souverainete europeenne, et il ne contredit pas la these : il existe au niveau groupe, pas Invent [D66], et il est non conclusif. Le groupe compte 2 partenaires souverains distincts (Mistral, Bleu) et 1 offre combinee (SAP-Bleu-Mistral) [D22, non verifie a la source primaire], contre 1 partenariat pour Accenture [D23] ; BCG X et Deloitte sont data not found ; McKinsey liste Mistral et Aleph Alpha dans son ecosysteme [D24, D63]. Ce n'est pas une avance etablie.")
para("Bilan par dimension (echelle 1-5, mediane M des 4 concurrents ; une dimension est non conclusive si plus de 2 acteurs sont en data not found, ND) :")
table([["Dimension", "Invent strict", "Groupe Capgemini", "Statut"],
       ["D1 Offres agentiques", "Retard provisoire (1 vs 3), du a l'absence d'attribution [D66]", "Parite (3 vs M = 3 ; parite avec McKinsey, 3 vs 3) [D01-D06]", "Conclusive au niveau groupe"],
       ["D2 Partenariats hyperscalers", "Non conclusive", "Non conclusive (indicatif 3, parite indicative avec Deloitte seulement) [D59, D62]", "ND : Accenture, McKinsey, BCG X, Invent [D61, D63, D64, D66]"],
       ["D3 Thought leadership (publications + conferences)", "Non conclusive", "Non conclusive", "ND pour les 5 acteurs en reprises presse [D26] ; conferences [D55-D58]"],
       ["D4 Taille des equipes IA", "Retard robuste vs Accenture (<= 3 vs 5) [D30, D35]", "Indetermine [D68]", "Non conclusive globalement"]], [3.4, 4.0, 4.2, 4.8], 8.5)
rich([("Compte : ", True), ("0 avance ; 0 retard ferme ; 1 retard robuste sur un sous-segment (D4 vs Accenture) ; 1 retard provisoire (D1 Invent strict) ; 1 parite (D1 groupe) ; 3 dimensions non conclusives. ", False),
      ("La cible du brief (2-3 avances, 2-3 retards) n'est pas atteinte et n'est pas forcee : lire des ND comme des scores violerait la regle du workplan.", True)])
rich([("Angle 2027 (hypothese a valider, pas recommandation ferme) : ", True), ("\"IA agentique souveraine et conforme en Europe, portee par Invent\". Elle exige d'abord de rattacher publiquement a Invent des actifs qui existent au niveau groupe. Elle est contestee par l'accord Accenture-Mistral de fevrier 2026 [S18] et par l'avance de McKinsey sur la remuneration au resultat [D42, non verifie a la source primaire].", False)])
rich([("4 decisions pour Invent (section 7) : ", True), ("(1) lancer un chantier d'attribution a 6 mois (baseline 0 actif IA attribue [D66], cible >= 2) ; (2) soumettre l'angle souverain agentique au test go/no-go (T1-T2 2027) ; (3) ne pas construire le positionnement sur l'echelle ni sur le seul outcome-based ; (4) relancer l'etude en mode standard (4 semaines) avant toute communication de positionnement.", False)])
d.add_page_break()

# 2 constat central
d.add_heading("2. Constat central : l'invisibilite d'Invent dans les sources publiques", 1)
para("Apres deux recherches ciblees, aucune source publique ne rattache a Invent une offre, un partenariat ou un effectif IA [D66, D69]. Les signaux de propriete pointent ailleurs :")
bl([("Etude agentique :", "publiee par le Capgemini Research Institute, pas par Invent [S66]."),
    ("Leadership agentique :", "affiche sous Insights & Data (\"Head of Agentic Transformation, Insights and Data\") [S67]."),
    ("Capital Markets Day 2026 :", "Invent n'apparait que comme une marque [D47, S4, verifie par fetch]."),
    ("Contraste :", "QuantumBlack et BCG X sont des marques IA autonomes, avec effectif [D33, D34], produits [D04, D05] et visibilite propres (QuantumBlack, Gold Partner de VivaTech 2025 [D55, S40]).")])
para("Consequence : Invent strict tombe au plancher des rubriques sans qu'une faiblesse ait ete mesuree. Pour un acheteur ou un analyste qui s'appuie sur des sources publiques, Invent n'a pas d'identite IA. Le rattachement est le prealable a tout angle de differenciation 2027.")

# 3 per dimension
d.add_heading("3. Detail par dimension", 1)
d.add_picture(f"{E}/03_charts/scoring_heatmap.png", width=Cm(16.4))
para("Figure 1. Scoring v2, echelle 1-5, mode express. Gris = ND ou NC ; ND ne signifie pas faible. Sources : scoring_matrix.json, D01-D69.", italic=True, size=8.5)
d.add_heading("3.1 D1 - Offres agentiques : un prerequis du marche, pas un differenciateur", 2)
para("Capgemini groupe obtient 3 [D01, D02 ; S1, S6], comme Accenture (AI Refinery for Industry, 12 solutions d'agents [D03, S21]), Deloitte (Zora AI, Global Agentic Network [D06, S30, S31]) et McKinsey (suite \"agents at scale\" [D04, S23]). BCG X est a 2 [D05]. Aucun acteur n'atteint le niveau 4, faute de references clients publiques comptees (DP02 : data not found [D07]).")
para("Preuves contraires : Accenture est probablement sous-note (references Ecolab/UOB non europeennes [D07, S15]) ; Accenture a cesse de publier ses metriques IA avancees apres le T1 FY26 [D09] ; le nombre d'agents McKinsey varie de 20 000 a 25 000 selon la source [D10].")
d.add_heading("3.2 D2 - Partenariats hyperscalers : non conclusive", 2)
para("Le top tier Capgemini n'est confirme que chez AWS [S53, verifie par fetch] et Google [S10] ; Microsoft est ND [D59, D60]. Deloitte egale Capgemini chez AWS [D62, S56]. Les tiers d'Accenture, McKinsey et BCG X ne sont pas publies, ou seulement sous forme d'awards (Accenture : 20 fois Microsoft Global SI Partner of the Year, 4 fois Google Services Partner of the Year [D61, S55]) ; McKinsey affiche des alliances avec les 3 hyperscalers sans termes [D63] ; BCG X [D64]. Les annuaires partenaires AWS, Microsoft et Google n'ont pas ete interroges [D65]. Accenture a le portefeuille de modeles de fondation le plus large (Anthropic, OpenAI, NVIDIA, Mistral) [D18].")
d.add_heading("3.3 D3 - Thought leadership (publications + conferences) : non conclusive", 2)
para("Les reprises presse sont ND pour les 5 acteurs [D26] ; les conferences sont ND ou indirectes pour 4 acteurs sur 5 [D55-D58]. Seul fait etabli : McKinsey QuantumBlack est Gold Partner de VivaTech 2025, et Capgemini est absent des tiers partenaires [D55, D56, S40, verifie par fetch]. Volumes de publications phares (bornes basses non comparables [D54]) : CRI 4-5, Deloitte 4, Accenture 3, McKinsey 1, BCG X 1 [D49-D53]. L'editeur est le CRI, pas Invent [S66]. A noter : PwC, EY et KPMG etaient sponsors de VivaTech 2025 [D55, S40], argument pour reintegrer PwC.")
d.add_heading("3.4 D4 - Taille des equipes IA : seul retard robuste, sur un sous-segment", 2)
table([["Acteur", "Effectif", "Niveau", "Source"],
       ["Accenture Data & AI", "77 000 professionnels AI & data (definition large)", "5", "D30, S15 (fetch)"],
       ["BCG X", "3 000+", "2", "D34"],
       ["QuantumBlack", "~1 700 (~7 000 selon le perimetre)", "2", "D33, S60"],
       ["Capgemini Invent", "\"10 000+\" toutes fonctions (texte de 2021) ; effectif IA non communique", "<= 3 (ND)", "D35, S11, D69"],
       ["Capgemini groupe", "Effectif data & IA non publie", "ND", "D68"],
       ["Deloitte", "ND", "ND", "-"]], [3.6, 7.0, 1.8, 4.0], 9)
para("Le retard d'Invent face a Accenture (<= 3 vs 5) tient quelle que soit la valeur reelle d'Invent. Face a QuantumBlack et BCG X, la comparaison reste ouverte (1-3 vs 2). Les 421 000 collaborateurs du groupe [D67] et les 310 000 formes a la GenAI [D38] ne sont pas des effectifs IA. Le comparatif melange une practice (Accenture), des entites (QuantumBlack, BCG X) et un think tank (Deloitte AI Institute). La repartition Europe est data not found [D37, DP14].")

# 4 hypotheses
d.add_heading("4. Verdict des 8 hypotheses du workplan", 1)
para("Les 8 hypotheses sont traitees. Aucune n'est confirmee, aucune n'est refutee : toutes restent sans reponse au sens strict de leurs criteres.")
table([["Hypothese", "Verdict (confiance)", "Motif principal"],
       ["H1.1 Parite avec McKinsey, retard sur Accenture/Deloitte en agentique", "Non testable (Low)", "Parite groupe avec McKinsey (3 vs 3) ; mais Accenture = Capgemini (3 vs 3) [D01, D03, D06] ; Invent non attribue [D66]."],
       ["H1.2 Avance sur l'agentique industriel en Europe", "Non testable", "Aucune reference agentique industrielle europeenne publique, pour aucun acteur [D14, DP04]."],
       ["H2.1 Parite hyperscalers avec Accenture/Deloitte, avance sur McKinsey", "Non testable ; signal faible contre l'avance (Low)", "Top tier confirme chez AWS et Google seulement ; Microsoft ND [D59, D60] ; Deloitte egale AWS [D62] ; tiers Accenture ND [D61]."],
       ["H2.2 Avance sur l'IA souveraine europeenne", "Non testable pour Invent ; indice favorable groupe (Low-Med)", "Groupe : 2 partenaires souverains distincts + 1 offre combinee [D22, non verifie a la source primaire] vs 1 Accenture [D23] ; McKinsey liste Mistral et Aleph Alpha (listing, pas accord) [D24, D63] ; rien attribue a Invent [D66] ; contre-indice S18."],
       ["H3.1 Retard sur McKinsey/BCG en visibilite, parite sur le volume", "Non testable (Low)", "Reprises presse ND [D26] ; conferences ND ou indirectes [D55-D58] ; volumes non comparables [D54]."],
       ["H4.1 Groupe 2e derriere Accenture ; Invent <= QuantumBlack et BCG X", "Non testable (Low)", "Effectifs groupe et Invent ND [D31, D68, D69] ; Invent vs QuantumBlack/BCG X indetermine [D33, D34]."],
       ["H5.1 Angle IA agentique industrielle et souveraine", "Non testable (Low)", "Depend de H1.2 et H2.2 ; demande documentee mais biaisee : 62 % [D45, S19, etude Accenture], 61 % [D46, S35, Gartner, non specifique a l'IA] ; DP17 : data not found."],
       ["H5.2 Angle agentique a l'echelle avec engagement de resultats", "Non testable (Low)", "Aucune offre Capgemini outcome-based annoncee [D41, S2, S4] ; McKinsey ~1/3 du revenu [D42, S37, non verifie a la source primaire] ; autres ND [D43, D44]."]], [5.4, 3.8, 7.2], 8.5)

# 5 angle
d.add_heading("5. Angle de differenciation 2027 : une hypothese a valider", 1)
rich([("Hypothese : ", True), ("\"Capgemini Invent, conseil de reference pour l'IA agentique souveraine et conforme en Europe\". Elle combine les actifs souverains du groupe (2 partenaires souverains distincts, Mistral et Bleu, plus 1 offre combinee SAP-Bleu-Mistral [D22 ; S7-S9, non verifie a la source primaire]), l'agentique NVIDIA/Google [D01, D02] et la demande souveraine [D45, D46]. Elle suppose que ces actifs soient explicitement attribues a Invent.", False)])
d.add_heading("Ce qui la soutient", 2)
bl(["Le groupe compte 2 partenaires souverains distincts (Mistral, Bleu) et 1 offre combinee (SAP-Bleu-Mistral) [D22, non verifie a la source primaire], contre 1 partenariat pour Accenture [D23] ; BCG X et Deloitte : data not found [D24].",
    "Le CMD 2026 nomme la souverainete parmi les marches en forte croissance [D47, S4].",
    "Demande documentee : 62 % des organisations europeennes cherchent une IA souveraine [D45] ; 61 % des DSI d'Europe de l'Ouest veulent plus de cloud local [D46] ; 48 % citent la conformite comme motivation principale [D45]."])
d.add_heading("Preuves contraires", 2)
bl(["Accenture-Mistral : accord multi-annuel, donnees hebergees dans l'UE, fevrier 2026 [D23, S18] ; Accenture publie aussi l'etude de reference sur la demande souveraine [S19].",
    "Le chiffre de 62 % vient d'une etude commanditee par un concurrent [D45] ; aucune etude independante ne mesure l'AI Act comme critere de choix (DP17 : data not found).",
    "McKinsey liste Mistral AI et Aleph Alpha dans son ecosysteme : un listing, pas un accord, mais les concurrents ne sont pas absents du terrain souverain [D24, D63].",
    "65 % des repondants disent ne pas pouvoir rivaliser sans fournisseurs non europeens ; seuls 19 % y voient un avantage competitif [D45].",
    "McKinsey tire environ un tiers de son revenu de modeles outcome-based [D42, S37, non verifie a la source primaire] et a lance son offre agentique a Paris (VivaTech 2025 [D04, D55]).",
    "L'actif est groupe, pas Invent [D66] : sans rattachement, l'angle profite a la marque Capgemini."])
d.add_heading("Test go / no-go", 2)
table([["#", "Test", "Seuil"],
       ["1", "Attribution", "Au moins 2 actifs souverains ou agentiques publiquement rattaches a Invent (capgemini.com/invent ou communique)"],
       ["2", "Preuve client", "Au moins 3 references clients europeennes publiques signees ou co-signees Invent, plus que chaque concurrent sur le segment (DP02/DP04)"],
       ["3", "Differenciation", "BCG X, McKinsey et Deloitte ont au plus 1 partenariat souverain europeen chacun (DP07) ; Accenture-Mistral ne couvre pas le meme perimetre de conseil"],
       ["4", "Demande", "Au moins 1 etude independante 2025-2026 (Commission europeenne, Eurostat, IDC/Gartner) montrant la souverainete ou l'AI Act comme critere de choix"]], [1, 3.2, 12.2], 9)
para("No-go si le test 1 echoue a 6 mois, ou si le test 3 montre au moins 2 concurrents avec au moins 2 partenariats souverains equivalents.")
table([["Element", "Contenu"],
       ["Horizon", "Validation T1-T2 2027 ; lancement au plus tot S2 2027 si go"],
       ["Owner", "Direction de Capgemini Invent (CEO Invent), avec le leader Data & AI groupe et le marketing d'offres Invent"],
       ["KPI", "(a) actifs IA attribues a Invent : 0 aujourd'hui [D66], cible >= 2 a 6 mois ; (b) references clients europeennes souveraines/agentiques publiques : cible >= 3 a 12 mois ; (c) part de voix Invent sur \"IA souveraine\" vs Accenture (mesure a construire, DP09 ND)"]], [3, 13.4], 9)

# 6 contraires
d.add_heading("6. Preuves contraires et divergences de donnees", 1)
bl([("Attribution :", "chaque actif Capgemini trouve releve du groupe ou du CRI ; aucun n'est rattache a Invent [D66 ; S2, S4]."),
    ("Hyperscalers :", "aucun avantage de tier etabli ; Deloitte a egalite chez AWS [D62] ; Accenture a le meilleur palmares d'awards [D61]."),
    ("Souverainete :", "l'avantage suppose du groupe s'erode (accord Accenture-Mistral [S18], etude Accenture [S19]) ; McKinsey liste Mistral et Aleph Alpha (listing, pas accord) [D24, D63] ; le comptage du groupe est 2 partenaires distincts + 1 offre combinee, pas 3 partenaires independants [D22, non verifie a la source primaire]."),
    ("Outcome-based :", "McKinsey ~1/3 du revenu [D42, non verifie a la source primaire] ; aucune offre Capgemini annoncee [D41]."),
    ("Chiffres divergents :", "agents McKinsey 20 000 vs 25 000 [D10] ; part IA de BCG 25 % vs > 40 % selon la definition [D11] ; QuantumBlack ~1 700 vs ~7 000 selon le perimetre [D33, S60]."),
    ("Comparabilite :", "practice (Accenture), entites (QuantumBlack, BCG X), think tank (Deloitte AI Institute).")])

# 7 decisions
d.add_heading("7. Les 4 decisions pour Invent", 1)
para("L'attribution passe en tete : sans rattachement des actifs groupe a Invent, aucun angle 2027 ne profite a Invent.")
table([["#", "Decision", "Horizon", "Owner", "KPI"],
       ["1", "Lancer un chantier d'attribution : rattacher publiquement a Invent des actifs IA du groupe (offre, partenariat ou etude)", "6 mois", "Direction de Capgemini Invent + marketing d'offres", "Actifs IA publiquement attribues a Invent : baseline 0 [D66], cible >= 2"],
       ["2", "Soumettre l'angle \"IA agentique souveraine europeenne\" au test go/no-go avant de l'adopter", "T1-T2 2027", "CEO d'Invent", "4 tests valides ; no-go si l'attribution echoue a 6 mois ou si >= 2 concurrents ont >= 2 partenariats souverains equivalents"],
       ["3", "Ne pas construire le positionnement 2027 sur l'echelle ni sur le seul outcome-based [D30 ; D42, non verifie a la source primaire]", "Immediat (cadrage du message 2027)", "Direction d'Invent", "Aucun claim d'echelle dans les supports Invent 2027"],
       ["4", "Relancer l'etude en mode standard avant toute communication de positionnement", "4 semaines", "Equipe engagement Atelier (sponsor : direction d'Invent)", "D2, D3, D4 passees de non conclusive a conclusive ; >= 80 % des S# cites verifies au primaire"]], [0.8, 6.0, 2.2, 3.2, 4.2], 8.5)

# 8 limites
d.add_heading("8. Limites de perimetre et plan de rerun en mode standard", 1)
bl([("Mode express :", "2 rerolls Hunter, budget data clos ; les scores sont un pre-brief, pas une evaluation finale."),
    ("Perimetre :", "PwC retire, D3 fusionnee, chiffres majoritairement groupe et mondial ; vue Europe quasi absente."),
    ("Qualite des sources :", "seules 8 sources sur 67 sont verifiees par fetch ; D22 (S7, S8, S9) et D42 (S37), qui portent l'angle 2027 et la decision 3, sont non verifies a la source primaire.")])
para("Lacunes a traiter en mode standard, par priorite (decision 4) :", bold=True)
for i, t in enumerate(["Rattachement a Invent : pages d'offres capgemini.com/invent, communiques, organigramme Data & AI [D66, D69].",
                       "Tiers hyperscalers via AWS Partner Finder, Microsoft Partner Finder et Google Cloud Partner Advantage, pour tous les acteurs [D65].",
                       "Effectifs data & IA : document d'enregistrement universel Capgemini, transcript du CMD, Deloitte AI & Data ; repartition Europe [D68, D69, D37].",
                       "References clients agentiques, au total, en Europe et dans l'industrie (DP02, DP04) ; partenariats souverains des concurrents (DP07), dont qualification des listings McKinsey [D63].",
                       "D3 : programmes speakers (VivaTech 2024/2026, Web Summit, London Tech Week), reprises presse (DP09) ; reintegration de PwC ; verification au primaire de S7, S8, S9 et S37."], 1):
    d.add_paragraph(t, style="List Number")

# annex methodo
d.add_heading("9. Annexes", 1)
d.add_heading("Annexe A. Methodologie et limites", 2)
bl([("Mode express :", "2 rerolls Hunter, budget data clos. Les scores sont un pre-brief, pas une evaluation finale."),
    ("Perimetre reduit :", "4 concurrents (McKinsey QuantumBlack, BCG X, Accenture, Deloitte). PwC AI Centre of Excellence retire (plafond express, perimetre public flou) ; EY, KPMG, IBM hors perimetre."),
    ("4 dimensions :", "publications et conferences fusionnees en D3, deux sous-scores lacunaires."),
    ("Notation :", "echelle 1-5 ; M = mediane des 4 concurrents ; non conclusive si plus de 2 acteurs en ND. Override D4 valide par le Chief (acteur focal ND, fourchette plausible 1-3)."),
    ("Qualite des sources :", "8 sources sur 67 verifiees par fetch (S2, S4, S15, S19, S40, S53, S55, S56) ; 59 viennent d'extraits de recherche et doivent etre reverifiees avant tout usage client."),
    ("Geographie :", "chiffres majoritairement groupe et mondial ; vue Europe quasi absente.")])

# bibliography
src = json.load(open(f"{E}/02_data/sources.json", encoding="utf8"))["sources"]
d.add_heading("Annexe B. Bibliographie (67 sources)", 2)
para("[v] = verifiee par fetch au primaire ; sinon extrait de recherche non verifie.", italic=True, size=9)
for s in src:
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"[{s['id']}] "); r.bold = True; r.font.size = Pt(8.5)
    tag = " [v]" if s.get("verified_by_fetch") else ""
    r = p.add_run(f"{s.get('publisher','')}, \"{s.get('title','')}\", {s.get('date','s.d.')}{tag}. {s.get('url','')}"); r.font.size = Pt(8.5)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Atelier | Pre-brief express | page "); run.font.size = Pt(8)
r2 = fp.add_run(); r2.font.size = Pt(8)
for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if t:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    r2._r.append(el)
d.save(f"{E}/06_deliverables/report.docx")
print("saved")
