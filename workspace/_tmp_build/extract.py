import sys, re
from pptx import Presentation
from pptx.util import Emu
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

E = sys.argv[1]
out = E + "/06_deliverables/"
ref = re.compile(r"\b(?:D\d+|S\d+|DP\d+|H\d\.\d)(?:-(?:D|S|DP)?\d+)?\b")

# ---- deck
prs = Presentation(out + "deck.pptx")
L = ["# Deck outline (extrait par relecture python-pptx de deck.pptx)", "",
     f"Nombre de slides : {len(prs.slides)}", ""]
for n, sl in enumerate(prs.slides, 1):
    texts = []; visuals = []; tables = []
    for sh in sl.shapes:
        if sh.shape_type == 13:
            visuals.append(f"image {sh.width/914400:.1f}x{sh.height/914400:.1f} in (scoring_heatmap.png)")
        elif sh.has_table if hasattr(sh, "has_table") else False:
            tables.append([[c.text for c in r.cells] for r in sh.table.rows])
        elif sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh.top, sh.left, sh.text_frame.text.strip()))
    texts.sort()
    title = None; source = None; body = []
    for top, left, t in texts:
        if n > 1 and title is None and top < 600000:
            title = t
        elif t.startswith("Source") and top > 6000000:
            source = t
        elif top > 6000000 and re.fullmatch(r"\d+", t):
            pass
        else:
            body.append(t)
    if n == 1:
        title = max((t for _, _, t in texts if len(t) > 60), key=lambda x: 0, default="")
        title = [t for top, l, t in texts if top > 800000 and top < 1000000][0]
        body = [t for top, l, t in texts if t != title and not t.startswith("Sources :")]
        source = [t for _, _, t in texts if t.startswith("Sources :")][0]
    L += [f"## Slide {n} - {title}", ""]
    L.append("**Contenu textuel**"); L.append("")
    for b in body:
        L.append("- " + b.replace("\n", "\n  "))
    for ti, tb_ in enumerate(tables, 1):
        L += ["", f"Tableau {ti} :", ""]
        L.append("| " + " | ".join(tb_[0]) + " |"); L.append("|" + "---|" * len(tb_[0]))
        for r in tb_[1:]:
            L.append("| " + " | ".join(c.replace("\n", " ").replace("|", "/") for c in r) + " |")
    alltext = " ".join([title] + body + [c for t in tables for r in t for c in r])
    cited = sorted(set(ref.findall(alltext + " " + (source or ""))))
    L += ["", f"**Sources citees (pied de slide)** : {source}", "",
          f"**References D#/S# relevees** : {', '.join(cited)}", "",
          f"**Visuel** : {'; '.join(visuals) if visuals else ('tableau' if tables else 'aucun (texte / encadres)')}", ""]
open(out + "deck_outline.md", "w", encoding="utf8").write("\n".join(L))

# ---- report
d = Document(out + "report.docx")
R = ["# Report outline (extrait par relecture python-docx de report.docx)", ""]
body = d.element.body
for ch in body.iterchildren():
    if ch.tag.endswith("}p"):
        p = Paragraph(ch, d)
        t = p.text.strip()
        if not t:
            if ch.xpath(".//w:drawing"):
                R.append("[Image inseree : scoring_heatmap.png]\n")
            continue
        st = p.style.name
        if st == "Heading 1": R.append("\n## " + t + "\n")
        elif st == "Heading 2": R.append("\n### " + t + "\n")
        elif st == "List Bullet": R.append("- " + t)
        elif st == "List Number": R.append("1. " + t)
        else: R.append(t + "\n")
    elif ch.tag.endswith("}tbl"):
        tb_ = Table(ch, d)
        rows = [[c.text.replace("\n", " ").replace("|", "/") for c in r.cells] for r in tb_.rows]
        if len(rows[0]) == 1:
            R.append("> " + rows[0][0] + "\n"); continue
        R.append("| " + " | ".join(rows[0]) + " |"); R.append("|" + "---|" * len(rows[0]))
        for r in rows[1:]: R.append("| " + " | ".join(r) + " |")
        R.append("")
open(out + "report_outline.md", "w", encoding="utf8").write("\n".join(R))
print("ok")
